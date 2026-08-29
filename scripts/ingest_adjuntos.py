import argparse
import json
import sys
import urllib.request
from pathlib import Path

import chromadb
from chromadb.utils.embedding_functions import OllamaEmbeddingFunction

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

TEXT_EXTENSIONS = {".txt", ".md", ".log", ".csv", ".json", ".xml", ".html", ".yml", ".yaml"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}


def login(backend_url, email, password):
    body = json.dumps({"email": email, "password": password}).encode("utf-8")
    req = urllib.request.Request(
        f"{backend_url}/api/auth/login",
        data=body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return data["token"]


def fetch_incidents(backend_url, token):
    req = urllib.request.Request(
        f"{backend_url}/api/incidents",
        headers={"Authorization": f"Bearer {token}"},
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def extraer_texto(path):
    ext = path.suffix.lower()
    if ext in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        try:
            import fitz
        except ImportError:
            print(f"  [warn] {path.name}: falta PyMuPDF para leer PDF")
            return ""
        doc = fitz.open(path)
        texto = "\n".join(page.get_text() for page in doc)
        doc.close()
        return texto
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            print(f"  [warn] {path.name}: falta python-docx para leer DOCX")
            return ""
        doc = Document(path)
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(cell.text for cell in fila.cells))
        return "\n".join(partes)
    if ext in IMAGE_EXTENSIONS:
        print(f"  [warn] {path.name}: tipo imagen, sin OCR en esta version")
        return ""
    print(f"  [warn] {path.name}: tipo {ext or 'desconocido'} no soportado")
    return ""


def split_document(text, chunk_size, overlap):
    chunks = []
    text = text.strip()
    if not text:
        return chunks
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end < len(text):
            break_at = text.rfind("\n", start, end)
            if break_at > start + chunk_size // 2:
                end = break_at
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = end - overlap
    return chunks


def main():
    parser = argparse.ArgumentParser(description="Ingiere los archivos adjuntos de los incidentes en chromadb (RAG local con Ollama).")
    parser.add_argument("--backend-url", default="http://localhost:8080",
                        help="URL del backend Spring (default: http://localhost:8080).")
    parser.add_argument("--admin-email", default="admin@helpdesk.com",
                        help="Email del admin para el GET /api/incidents (default: admin@helpdesk.com).")
    parser.add_argument("--admin-password", default="admin123",
                        help="Password del admin (default: admin123).")
    parser.add_argument("--uploads-dir", default="uploads",
                        help="Carpeta con los archivos adjuntos (default: uploads).")
    parser.add_argument("--embed-model", default="nomic-embed-text",
                        help="Modelo de Ollama para embeddings (default: nomic-embed-text).")
    parser.add_argument("--ollama-url", default="http://localhost:11434",
                        help="URL base del servidor Ollama (default: http://localhost:11434).")
    parser.add_argument("--chunk-size", type=int, default=800,
                        help="Tamano de cada chunk en caracteres (default: 800).")
    parser.add_argument("--overlap", type=int, default=100,
                        help="Superposicion entre chunks en caracteres (default: 100).")
    parser.add_argument("--db", default="kb", help="Carpeta persistente de chromadb (default: kb).")
    parser.add_argument("--collection", default="adjuntos",
                        help="Nombre de la coleccion (default: adjuntos).")
    parser.add_argument("--append", action="store_true",
                        help="No borrar la coleccion existente antes de ingerir (por defecto se reconstruye).")
    args = parser.parse_args()

    root = Path(__file__).resolve().parent.parent
    uploads_dir = Path(args.uploads_dir)
    if not uploads_dir.is_absolute():
        uploads_dir = root / uploads_dir

    token = login(args.backend_url, args.admin_email, args.admin_password)
    incidentes = fetch_incidents(args.backend_url, token)

    adjuntos = {}
    for inc in incidentes:
        if inc.get("archivoUrl"):
            adjuntos[inc["archivoUrl"]] = inc

    if not adjuntos:
        print(f"[ingest-adjuntos] No hay incidentes con archivos adjuntos.")
        return

    client = chromadb.PersistentClient(path=args.db)
    if not args.append:
        try:
            client.delete_collection(args.collection)
            print(f"[ingest-adjuntos] Coleccion '{args.collection}' eliminada (reconstruccion).")
        except Exception:
            pass

    embedding_fn = OllamaEmbeddingFunction(url=args.ollama_url, model_name=args.embed_model)
    collection = client.get_or_create_collection(args.collection, embedding_function=embedding_fn)

    total = 0
    for filename, inc in sorted(adjuntos.items()):
        path = uploads_dir / filename
        if not path.exists():
            print(f"  [warn] {filename}: no existe en {uploads_dir}")
            continue
        texto = extraer_texto(path)
        chunks = split_document(texto, args.chunk_size, args.overlap)
        if not chunks:
            continue
        etiqueta = f"Adjunto #{inc['id']} {inc.get('titulo', '')}".strip()
        ids = [f"adjunto:{filename}#{i}" for i in range(len(chunks))]
        metadatas = [{
            "archivo": etiqueta,
            "filename": filename,
            "incident_id": inc["id"],
            "incident_titulo": inc.get("titulo", ""),
            "tipo": "adjunto",
        } for _ in range(len(chunks))]
        collection.add(ids=ids, documents=chunks, metadatas=metadatas)
        total += len(chunks)
        print(f"[ingest-adjuntos] {etiqueta} ({filename}): {len(chunks)} chunks")

    print(f"[ingest-adjuntos] Listo. {total} chunks en la coleccion '{args.collection}' (embedding: {args.embed_model}).")


if __name__ == "__main__":
    main()