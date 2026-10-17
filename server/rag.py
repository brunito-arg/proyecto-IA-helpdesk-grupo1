import os
from pathlib import Path

import chromadb
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

TEXT_EXTENSIONS = {".txt", ".md", ".log", ".csv", ".json", ".xml", ".html", ".yml", ".yaml"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}


def ollama_url(base_url=None):
    return base_url or os.getenv("OLLAMA_URL", "http://localhost:11434")


def embed_model_name(model=None):
    return model or os.getenv("EMBED_MODEL", "nomic-embed-text")


def make_embeddings(model=None, base_url=None):
    return OllamaEmbeddings(model=embed_model_name(model), base_url=ollama_url(base_url))


def make_splitter(chunk_size=800, chunk_overlap=100):
    return RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )


def split_document(text, metadata, chunk_size=800, chunk_overlap=100, id_prefix=None):
    if not text or not text.strip():
        return [], []
    chunks = [c for c in make_splitter(chunk_size, chunk_overlap).split_text(text) if c.strip()]
    documentos = []
    ids = []
    for i, chunk in enumerate(chunks):
        meta = dict(metadata)
        meta["indice"] = i
        documentos.append(Document(page_content=chunk, metadata=meta))
        ids.append(f"{id_prefix}#{i}" if id_prefix is not None else str(i))
    return documentos, ids


def get_vectorstore(collection, persist_dir, model=None, base_url=None, client=None):
    kwargs = {
        "collection_name": collection,
        "embedding_function": make_embeddings(model, base_url),
    }
    if client is not None:
        kwargs["client"] = client
    else:
        kwargs["persist_directory"] = str(persist_dir)
    return Chroma(**kwargs)


def open_collections(names, persist_dir, model=None, base_url=None):
    client = chromadb.PersistentClient(path=str(persist_dir))
    stores = []
    for name in names:
        try:
            store = get_vectorstore(name, persist_dir, model, base_url, client=client)
            total = store._collection.count()
            if total == 0:
                print(f"[rag] Coleccion '{name}' vacia; corré la ingesta correspondiente.")
                continue
            print(f"[rag] Base de conocimiento cargada: '{name}' ({total} chunks)")
            stores.append(store)
        except Exception as e:
            print(f"[rag] No se pudo abrir la coleccion '{name}': {e}")
    return stores


def recuperar_contexto(pregunta, stores, top_k=5):
    if not stores:
        return ""
    por_coleccion = max(1, top_k // max(1, len(stores)))
    fragmentos = []
    for store in stores:
        try:
            fragmentos.extend(store.similarity_search(pregunta, k=por_coleccion))
        except Exception as e:
            nombre = getattr(store._collection, "name", "?")
            raise RuntimeError(f"Error al buscar en la coleccion '{nombre}': {e}")
    vistos = set()
    unicos = []
    for doc in fragmentos:
        if doc.page_content in vistos:
            continue
        vistos.add(doc.page_content)
        unicos.append(doc)
    return "\n\n".join(
        f"[Fuente: {d.metadata.get('archivo', '?')}] {d.page_content}" for d in unicos[:top_k]
    )


def extraer_texto(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext in TEXT_EXTENSIONS:
        try:
            from langchain_community.document_loaders import TextLoader

            loader = TextLoader(str(path), encoding="utf-8", autodetect_encoding=True)
            return "\n".join(d.page_content for d in loader.load())
        except Exception:
            return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        try:
            from langchain_community.document_loaders import PyMuPDFLoader

            return "\n".join(d.page_content for d in PyMuPDFLoader(str(path)).load())
        except ImportError:
            import fitz

            doc = fitz.open(path)
            texto = "\n".join(page.get_text() for page in doc)
            doc.close()
            return texto
    if ext == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(cell.text for cell in fila.cells))
        return "\n".join(partes)
    if ext in IMAGE_EXTENSIONS:
        return ""
    return ""
